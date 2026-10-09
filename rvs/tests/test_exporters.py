"""Structure and determinism of every output format (golden files are in test_golden.py)."""

import io
import re
from datetime import datetime
from pathlib import Path
from typing import Any

import pytest

from rvs_core.adapter import DoorstopProject
from rvs_core.config import load_project_config
from rvs_core.exporters import DOC_FORMATS, TABLE_FORMATS, render_doc, render_table
from rvs_core.exporters.builders import matrix_doc, spec_doc
from rvs_core.matrices import Provenance, build_traceability, build_vcm
from rvs_core.trace import LinkGraph

PROV = Provenance("0.1.0", "3.2", "Minimal example", "working copy", datetime(2026, 1, 2, 3, 4, 5), "alice")


def _ctx(root: Path) -> tuple[Any, Any, Any]:
    cfg, _ = load_project_config(root)
    items = DoorstopProject.open(root).items()
    return cfg, items, LinkGraph.build(cfg, items)


@pytest.fixture
def vcm(minimal_project: Path):  # type: ignore[no-untyped-def]
    cfg, items, graph = _ctx(minimal_project)
    return build_vcm(cfg, items, graph, provenance=PROV)


@pytest.fixture
def spec(minimal_project: Path):  # type: ignore[no-untyped-def]
    cfg, items, graph = _ctx(minimal_project)
    return spec_doc(cfg, items, graph, ["SYS", "EPS"], PROV)


# builders -------------------------------------------------------------------------------
def test_spec_doc_contains_entries_headings_and_links(spec):  # type: ignore[no-untyped-def]
    from rvs_core.exporters.model import Entry, Heading

    entries = [b for b in spec.blocks if isinstance(b, Entry)]
    assert [e.uid for e in entries] == ["SYS-0001", "SYS-0002", "SYS-0003", "EPS-0001", "EPS-0002", "EPS-0003"]
    heads = [b.text for b in spec.blocks if isinstance(b, Heading)]
    assert (
        "SYS — System requirements" in heads and "System requirements" in heads
    )  # document + non-normative heading item
    eps1 = next(e for e in entries if e.uid == "EPS-0001")
    fields = dict(eps1.fields)
    assert fields["Parents"] == "SYS-0002" and fields["Verified by"] == "VER-0001" and fields["Status"] == "draft"
    assert "shall store" in eps1.text


def test_spec_doc_rejects_unknown_document(minimal_project: Path):
    cfg, items, graph = _ctx(minimal_project)
    with pytest.raises(ValueError, match="NOPE"):
        spec_doc(cfg, items, graph, ["NOPE"], PROV)


def test_matrix_doc_wraps_table_and_notes(vcm):  # type: ignore[no-untyped-def]
    doc = matrix_doc(vcm)
    assert doc.title == "Verification control matrix" and doc.notes and doc.provenance == PROV


# html ------------------------------------------------------------------------------------
def test_html_is_self_contained_and_escaped(spec):  # type: ignore[no-untyped-def]
    html = render_doc(spec, "html").decode("utf-8")
    assert html.startswith("<!doctype html>") and "System requirements" in html
    assert "Generated: 2026-01-02 03:04:05 by alice" in html and "Project: Minimal example" in html
    assert "data:font/woff;base64," in html and "@font-face" in html
    assert not re.search(r"(src|href)=[\"']https?://", html) and "<script" not in html
    assert 'id="EPS-0001"' in html


def test_html_escapes_markup_in_statements(minimal_project: Path):
    proj = DoorstopProject.open(minimal_project)
    proj.update_item("SYS-0001", text="The system shall handle <b>raw</b> & <script>x()</script> input.")
    cfg, items, graph = _ctx(minimal_project)
    html = render_doc(spec_doc(cfg, items, graph, ["SYS"], PROV), "html").decode("utf-8")
    assert "<script>" not in html and "&lt;script&gt;" in html and "&amp;" in html


def test_html_matrix_marks_gap_rows(minimal_project: Path):
    cfg, items, graph = _ctx(minimal_project)
    t = build_traceability(cfg, items, graph, "SYS", "VER", "down", provenance=PROV)
    html = render_table(t, "html").decode("utf-8")
    assert 'class="gap-unverified"' in html


# docx --------------------------------------------------------------------------------------
def test_docx_opens_and_has_table_title_block_and_properties(vcm):  # type: ignore[no-untyped-def]
    from docx import Document

    data = render_table(vcm, "docx")
    doc = Document(io.BytesIO(data))
    assert doc.core_properties.title == "Verification control matrix"
    assert doc.core_properties.author == "alice" and doc.core_properties.created.replace(tzinfo=None) == datetime(
        2026, 1, 2, 3, 4, 5
    )
    table = doc.tables[0]
    assert [c.text for c in table.rows[0].cells][0] == "Requirement" and len(table.rows) == 1 + 6
    header_text = " ".join(p.text for p in doc.sections[0].header.paragraphs)
    assert "Minimal example" in header_text and "alice" in header_text
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "TODO-STANDARD" in body


def test_docx_spec_has_entries(spec):  # type: ignore[no-untyped-def]
    from docx import Document

    doc = Document(io.BytesIO(render_doc(spec, "docx")))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "EPS-0001" in text and "shall store at least the energy" in text and "Parents: SYS-0002" in text


# pdf -----------------------------------------------------------------------------------------
def test_pdf_has_text_fonts_and_provenance(spec):  # type: ignore[no-untyped-def]
    from pypdf import PdfReader

    data = render_doc(spec, "pdf")
    assert data.startswith(b"%PDF-")
    reader = PdfReader(io.BytesIO(data))
    text = "\n".join(p.extract_text() for p in reader.pages)
    assert "EPS-0001" in text and "Minimal example" in text and "alice" in text
    fonts = {str(f.get("/BaseFont", "")) for p in reader.pages for f in (p["/Resources"].get("/Font") or {}).values()
             for f in [f.get_object()]}  # fmt: skip
    assert any("IBMPlexSans" in f for f in fonts)  # bundled font, embedded (not Helvetica)


def test_pdf_matrix_with_unicode_and_wide_table(vcm):  # type: ignore[no-untyped-def]
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(render_table(vcm, "pdf")))
    text = "\n".join(p.extract_text() for p in reader.pages)
    assert "Requirement" in text and "EPS-0001" in text
    assert float(reader.pages[0].mediabox.width) > float(reader.pages[0].mediabox.height)  # wide table: landscape


def test_pdf_handles_non_latin_text(minimal_project: Path):
    from pypdf import PdfReader

    DoorstopProject.open(minimal_project).update_item(
        "SYS-0001", text="The system shall keep 5 µm ≤ gap → 10 Ω ✓ 温度."
    )
    cfg, items, graph = _ctx(minimal_project)
    data = render_doc(spec_doc(cfg, items, graph, ["SYS"], PROV), "pdf")
    assert PdfReader(io.BytesIO(data)).pages  # unsupported glyphs must not crash the export


# xlsx -------------------------------------------------------------------------------------------
def test_xlsx_matrix(vcm):  # type: ignore[no-untyped-def]
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(render_table(vcm, "xlsx")))
    assert wb.sheetnames == ["Matrix", "Provenance"]
    ws = wb["Matrix"]
    assert [c.value for c in ws[1]][:2] == ["Requirement", "Title"] and ws.max_row == 7
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref
    prov = {r[0].value: r[1].value for r in wb["Provenance"].iter_rows()}
    assert prov["Project"] == "Minimal example" and prov["Generated by"] == "alice"
    assert wb.properties.creator == "alice"


def test_xlsx_never_creates_formulas(minimal_project: Path):
    from openpyxl import load_workbook

    DoorstopProject.open(minimal_project).update_item("SYS-0001", attrs={"title": '=HYPERLINK("http://x","y")'})
    cfg, items, graph = _ctx(minimal_project)
    wb = load_workbook(io.BytesIO(render_table(build_vcm(cfg, items, graph, provenance=PROV), "xlsx")))
    cell = wb["Matrix"]["B2"]
    assert cell.data_type == "s" and cell.value.startswith("=HYPERLINK")


# csv / json -------------------------------------------------------------------------------------
def test_csv_and_json_bytes(vcm):  # type: ignore[no-untyped-def]
    assert render_table(vcm, "csv").decode("utf-8").startswith("# Verification control matrix")
    assert render_table(vcm, "json").decode("utf-8").lstrip().startswith("{")


# determinism ---------------------------------------------------------------------------------------
@pytest.mark.parametrize("fmt", TABLE_FORMATS)
def test_every_table_format_is_byte_identical_for_identical_input(vcm, fmt: str):  # type: ignore[no-untyped-def]
    assert render_table(vcm, fmt) == render_table(vcm, fmt)


@pytest.mark.parametrize("fmt", DOC_FORMATS)
def test_every_doc_format_is_byte_identical_for_identical_input(spec, fmt: str):  # type: ignore[no-untyped-def]
    assert render_doc(spec, fmt) == render_doc(spec, fmt)


def test_unknown_format_is_an_error(vcm, spec):  # type: ignore[no-untyped-def]
    with pytest.raises(ValueError, match="rtf"):
        render_table(vcm, "rtf")
    with pytest.raises(ValueError, match="xlsx"):
        render_doc(spec, "xlsx")
