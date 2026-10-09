import subprocess
import sys
from pathlib import Path

from rvs_cli.main import main


def _export(minimal_project: Path, tmp_path: Path, *args: str, name: str = "out") -> Path:
    target = tmp_path / name
    assert main(["export", str(minimal_project), *args, "--output", str(target)]) == 0
    return target


def test_binary_formats_need_an_output_file(minimal_project: Path):
    out = subprocess.run(
        [sys.executable, "-m", "rvs_cli", "export", str(minimal_project), "--vcm", "--format", "pdf"],
        capture_output=True,
        check=False,
    )
    assert out.returncode == 2 and b"--output" in out.stderr and out.stdout == b""


def test_vcm_in_every_format(minimal_project: Path, tmp_path: Path):
    from docx import Document
    from openpyxl import load_workbook
    from pypdf import PdfReader

    xlsx = _export(minimal_project, tmp_path, "--vcm", "--format", "xlsx", name="v.xlsx")
    assert "Matrix" in load_workbook(xlsx).sheetnames
    assert Document(str(_export(minimal_project, tmp_path, "--vcm", "--format", "docx", name="v.docx"))).tables
    assert PdfReader(str(_export(minimal_project, tmp_path, "--vcm", "--format", "pdf", name="v.pdf"))).pages
    assert "<table>" in _export(minimal_project, tmp_path, "--vcm", "--format", "html", name="v.html").read_text(
        "utf-8"
    )


def test_items_export_csv_and_xlsx(minimal_project: Path, tmp_path: Path):
    csv_file = _export(minimal_project, tmp_path, "--items", name="items.csv")
    assert csv_file.read_bytes().startswith(b"\xef\xbb\xbf")
    xlsx_file = _export(minimal_project, tmp_path, "--items", "--format", "xlsx", name="items.xlsx")
    from openpyxl import load_workbook

    assert load_workbook(xlsx_file).sheetnames == ["SYS", "EPS", "VER", "Provenance"]


def test_items_export_rejects_document_formats(minimal_project: Path):
    assert main(["export", str(minimal_project), "--items", "--format", "pdf", "-o", "x.pdf"]) == 2


def test_spec_export(minimal_project: Path, tmp_path: Path):
    from pypdf import PdfReader

    pdf = _export(minimal_project, tmp_path, "--spec", "--document", "EPS", "--format", "pdf", name="s.pdf")
    text = "\n".join(p.extract_text() for p in PdfReader(str(pdf)).pages)
    assert "EPS-0001" in text and "System requirements" not in text
    assert "EPS-0001" in _export(minimal_project, tmp_path, "--spec", "--format", "html", name="s.html").read_text(
        "utf-8"
    )
    assert main(["export", str(minimal_project), "--spec", "--format", "csv", "-o", str(tmp_path / "x.csv")]) == 2


def test_import_dry_run_reports_and_changes_nothing(minimal_project: Path, tmp_path: Path, capsys):  # type: ignore[no-untyped-def]
    csv_file = _export(minimal_project, tmp_path, "--items", name="items.csv")
    text = csv_file.read_text("utf-8-sig").replace("Payload power", "Payload power v2")
    csv_file.write_text(text, encoding="utf-8")
    before = (minimal_project / "SYS" / "SYS-0001.yml").read_bytes()
    capsys.readouterr()
    code = main(["import", str(minimal_project), str(csv_file), "--dry-run"])
    out = capsys.readouterr().out
    assert code == 0 and "update" in out and "SYS-0001" in out and "title" in out and "dry run" in out.lower()
    assert (minimal_project / "SYS" / "SYS-0001.yml").read_bytes() == before


def test_import_applies_and_records_the_reason(minimal_project: Path, tmp_path: Path, capsys):  # type: ignore[no-untyped-def]
    from rvs_core.authoring import read_history

    csv_file = _export(minimal_project, tmp_path, "--items", name="items.csv")
    csv_file.write_text(csv_file.read_text("utf-8-sig").replace("Payload power", "Payload power v2"), encoding="utf-8")
    code = main(["import", str(minimal_project), str(csv_file), "--reason", "customer review", "--user", "carol"])
    assert code == 0 and "1 updated" in capsys.readouterr().out
    entry = read_history(minimal_project, "SYS-0001")[-1]
    assert entry["who"] == "carol" and entry["why"] == "customer review" and entry["action"] == "import"


def test_import_xlsx_round_trip_is_a_no_op(minimal_project: Path, tmp_path: Path, capsys):  # type: ignore[no-untyped-def]
    xlsx = _export(minimal_project, tmp_path, "--items", "--format", "xlsx", name="items.xlsx")
    capsys.readouterr()
    assert main(["import", str(minimal_project), str(xlsx)]) == 0
    assert "10 unchanged" in capsys.readouterr().out


def test_import_errors_exit_one_and_apply_nothing(minimal_project: Path, tmp_path: Path, capsys):  # type: ignore[no-untyped-def]
    bad = tmp_path / "bad.csv"
    bad.write_text("id,status\nSYS-0001,finished\n", encoding="utf-8")
    code = main(["import", str(minimal_project), str(bad)])
    out = capsys.readouterr().out
    assert code == 1 and "ERROR" in out and "finished" in out and "nothing was imported" in out.lower()


def test_import_usage_and_load_errors(minimal_project: Path, tmp_path: Path, capsys):  # type: ignore[no-untyped-def]
    assert main(["import", str(minimal_project), str(tmp_path / "missing.csv")]) == 2
    odd = tmp_path / "x.txt"
    odd.write_text("a")
    assert main(["import", str(minimal_project), str(odd)]) == 2
    csv_file = tmp_path / "ok.csv"
    csv_file.write_text("id\n")
    assert main(["import", str(tmp_path), str(csv_file)]) == 3
