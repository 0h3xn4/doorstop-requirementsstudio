"""Item tables: export to CSV/XLSX and import from them.

One row per item. Columns are machine names (``id, document, level, normative, derived, active, header, ref, text,
parents`` and every attribute of the document templates). On import only the columns present are considered, so a
re-imported export changes nothing, and a file with just ``id,title`` updates just titles. Imports are planned first
(dry run) and applied only when the plan is clean, or with ``skip_errors``.
"""

import csv
import io
import re
import sys
import zipfile
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment

from rvs_core import csvsafe, textcheck
from rvs_core.adapter import DoorstopProject, ItemData, ProjectError
from rvs_core.authoring import EditService
from rvs_core.changecontrol.manifests import baselined_uids
from rvs_core.config import ProjectConfig
from rvs_core.config.model import AttributeDef
from rvs_core.exporters.xlsx_out import XLSX_CELL_LIMIT, finish, fit_columns, provenance_sheet, set_text, style_header
from rvs_core.matrices.provenance import Provenance

CORE_COLUMNS = ["id", "document", "level", "normative", "derived", "active", "header", "ref", "text", "parents"]
ALIASES = {"uid": "id", "prefix": "document", "statement": "text", "parent": "parents"}
HIDDEN_ATTRIBUTES = {"rvs_schema_version"}
_UID = re.compile(r"^[A-Z][A-Z0-9]*-\d+$")
_LEVEL = re.compile(r"^\d+(\.\d+)*$")
_DATE = re.compile(r"^[0-9]{4}-[0-9]{2}-[0-9]{2}$")
_TRUE, _FALSE = {"yes", "true", "1", "y"}, {"no", "false", "0", "n"}
csv.field_size_limit(min(sys.maxsize, 2**31 - 1))  # RVS can write a statement larger than the default 128 KiB limit
MAX_XLSX_ROWS, MAX_XLSX_COLUMNS, MAX_XLSX_BYTES = 200_000, 300, 400 * 1024 * 1024


# columns and value encoding ###################################################################
def item_columns(cfg: ProjectConfig) -> list[str]:
    cols = list(CORE_COLUMNS)
    for kind in ("requirements", "verification"):
        for a in cfg.templates.kinds[kind].attributes:
            if a.name not in HIDDEN_ATTRIBUTES and a.name not in cols:
                cols.append(a.name)
    for a in cfg.project.free_attributes:
        if a.name not in cols:
            cols.append(a.name)
    return cols


def _attr_defs(cfg: ProjectConfig) -> dict[str, dict[str, AttributeDef]]:
    return {k: cfg.attribute_defs(k) for k in ("requirements", "verification")}


def _encode_ref_list(value: Any) -> str:
    return "; ".join(f"{v['docno']}:{v['revision']}" for v in value or [])


def encode(value: Any, adef: AttributeDef | None = None) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    if adef is not None and adef.type == "ref-list":
        return _encode_ref_list(value)
    if adef is not None and adef.type == "string-list" and isinstance(value, list | tuple):
        return ", ".join(_escape(str(v)) for v in value)
    if isinstance(value, list | tuple):
        return ", ".join(str(v) for v in value)
    return str(value)


def _row_for(
    cfg: ProjectConfig, item: ItemData, columns: Sequence[str], defs: dict[str, dict[str, AttributeDef]]
) -> list[str]:
    kind = (cfg.project.document(item.document) or None) and cfg.project.document(item.document).kind  # type: ignore[union-attr]
    core = {
        "id": item.uid, "document": item.document, "level": item.level, "normative": encode(item.normative),
        "derived": encode(item.derived), "active": encode(item.active), "header": item.header, "ref": item.ref,
        "text": item.text.strip(), "parents": ", ".join(item.links),
    }  # fmt: skip
    out = []
    for c in columns:
        if c in core:
            out.append(core[c])
        else:
            adef = defs.get(kind or "", {}).get(c)
            out.append(encode(item.attrs.get(c), adef) if adef else "")
    return out


def _ordered(cfg: ProjectConfig, items: Sequence[ItemData]) -> list[ItemData]:
    order = {d.prefix: n for n, d in enumerate(cfg.project.documents)}
    return sorted(items, key=lambda i: (order.get(i.document, 999), i.level_key, i.uid))


def export_items_csv(cfg: ProjectConfig, items: Sequence[ItemData], prov: Provenance) -> bytes:
    cols, defs = item_columns(cfg), _attr_defs(cfg)
    buf = io.StringIO()
    buf.write("# Items\n")
    for line in prov.lines():
        buf.write(f"# {line}\n")
    writer = csv.writer(buf, lineterminator="\n")
    writer.writerow(cols)
    writer.writerows([csvsafe.protect(c) for c in _row_for(cfg, i, cols, defs)] for i in _ordered(cfg, items))
    return b"\xef\xbb\xbf" + buf.getvalue().encode("utf-8")  # BOM: Excel opens UTF-8 CSV correctly


def export_items_xlsx(cfg: ProjectConfig, items: Sequence[ItemData], prov: Provenance) -> bytes:
    cols, defs = item_columns(cfg), _attr_defs(cfg)
    wb = Workbook()
    wb.remove(wb.active)
    ordered = _ordered(cfg, items)
    wrap = Alignment(vertical="top", wrap_text=True)
    for decl in cfg.project.documents:
        ws = wb.create_sheet(decl.prefix)
        rows = [_row_for(cfg, i, cols, defs) for i in ordered if i.document == decl.prefix]
        for row in rows:
            for column, cell in zip(cols, row, strict=True):
                if len(cell) > XLSX_CELL_LIMIT:
                    raise ValueError(
                        f"{row[0]}: the {column} is {len(cell):,} characters long, more than an Excel cell holds "
                        f"({XLSX_CELL_LIMIT:,}). Export the items as CSV or ReqIF instead."
                    )
        for c, name in enumerate(cols, start=1):
            ws.cell(row=1, column=c, value=name)
        for r, row in enumerate(rows, start=2):
            for c, value in enumerate(row, start=1):
                cell = ws.cell(row=r, column=c)
                set_text(cell, value)
                cell.alignment = wrap
        style_header(ws, len(cols))
        fit_columns(ws, cols, rows)
    provenance_sheet(wb, prov)
    return finish(wb, prov, "Items")


# reading ##########################################################################################
def _canonical_header(names: Sequence[str]) -> list[str]:
    return [ALIASES.get(n.strip().lower(), n.strip().lower()) for n in names]


def read_csv(data: bytes) -> list[dict[str, str]]:
    text = data.decode("utf-8-sig")
    lines = text.splitlines(keepends=True)
    start = 0
    while start < len(lines) and lines[start].lstrip().startswith("#"):
        start += 1  # leading provenance/comment lines
    reader = csv.reader(io.StringIO("".join(lines[start:]), newline=""))
    try:
        header = _canonical_header(next(reader))
    except StopIteration:
        return []
    rows = []
    for record in reader:
        if not any(c.strip() for c in record):
            continue
        row = {
            h: (csvsafe.unprotect(record[i]).replace("\r\n", "\n").replace("\r", "\n") if i < len(record) else "")
            for i, h in enumerate(header)
        }
        row["_row"] = str(start + reader.line_num)
        extra = [c for c in record[len(header) :] if c.strip()]
        if extra:
            row["_error"] = (
                f"The row has {len(extra)} more cell(s) than the header has columns. "
                "An unquoted comma inside a text is the usual cause: put that text in quotes"
            )
        rows.append(row)
    return rows


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    return str(value).replace("\r\n", "\n").replace("\r", "\n")


def read_xlsx(data: bytes) -> list[dict[str, str]]:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_XLSX_BYTES:
                raise ValueError("The workbook expands to more than 400 MB, which RVS will not read")
    except zipfile.BadZipFile:
        raise ValueError("The file is not an Excel workbook (.xlsx)") from None
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)  # streams; a sparse sheet costs nothing
    rows: list[dict[str, str]] = []
    try:
        for ws in wb.worksheets:
            if ws.title.lower() in ("provenance", "readme"):
                continue
            it = ws.iter_rows(min_row=1, max_row=MAX_XLSX_ROWS + 1, max_col=MAX_XLSX_COLUMNS, values_only=True)
            try:
                header = _canonical_header([_cell_text(v) for v in next(it)])
            except StopIteration:
                continue
            empty_run = 0
            for n, values in enumerate(it, start=2):
                cells = [_cell_text(v) for v in values]
                if not any(c.strip() for c in cells):
                    empty_run += 1
                    if empty_run >= 1000:
                        break  # the rest of a sparse sheet is empty
                    continue
                empty_run = 0
                if n > MAX_XLSX_ROWS:
                    raise ValueError(f"The sheet {ws.title} has more than {MAX_XLSX_ROWS:,} rows")
                row = {h: (cells[i] if i < len(cells) else "") for i, h in enumerate(header) if h}
                if not row.get("document", "").strip():
                    row["document"] = ws.title
                row["_row"] = str(n)
                row["_sheet"] = ws.title
                rows.append(row)
    finally:
        wb.close()
    return rows


# planning ###########################################################################################
@dataclass(frozen=True)
class RowResult:
    row: int
    uid: str
    action: str  # create | update | unchanged | error
    changes: tuple[str, ...] = ()
    message: str = ""
    document: str = ""


@dataclass
class _Op:
    result_index: int
    kind: str
    document: str
    uid: str
    number: int | None
    core: dict[str, Any]
    attrs: dict[str, Any]
    parents: list[str] | None
    changes: list[str]


@dataclass
class ImportPlan:
    results: list[RowResult] = field(default_factory=list)
    ops: list[_Op] = field(default_factory=list)

    @property
    def errors(self) -> list[RowResult]:
        return [r for r in self.results if r.action == "error"]

    def count(self, action: str) -> int:
        return sum(1 for r in self.results if r.action == action)


@dataclass(frozen=True)
class ImportReport:
    applied: bool
    created: int = 0
    updated: int = 0
    unchanged: int = 0
    errors: int = 0
    created_uids: tuple[str, ...] = ()


def _bool(text: str) -> bool:
    t = text.strip().lower()
    if t in _TRUE:
        return True
    if t in _FALSE:
        return False
    raise ValueError(f"'{text}' is not yes or no")


def _escape(value: str) -> str:
    return value.replace("\\", "\\\\").replace(",", "\\,").replace(";", "\\;")


def _split_escaped(text: str) -> list[str]:
    """Split a string-list cell on , or ; where they are not escaped with a backslash."""
    parts, current, i = [], [], 0
    while i < len(text):
        c = text[i]
        if c == "\\" and i + 1 < len(text):
            current.append(text[i + 1])
            i += 2
            continue
        if c in ",;":
            parts.append("".join(current).strip())
            current = []
        else:
            current.append(c)
        i += 1
    parts.append("".join(current).strip())
    return [p for p in parts if p]


def _split(text: str) -> list[str]:
    return [p.strip() for p in re.split(r"[,;]", text) if p.strip()]


def _decode_attr(adef: AttributeDef, text: str, cfg: ProjectConfig) -> Any:
    t = adef.type
    if t in ("string", "text"):
        return text
    if t == "enum":
        value = text.strip()
        allowed = cfg.vocab.values(adef.vocab or "")
        if value and value not in allowed:
            raise ValueError(f"'{value}' is not allowed for {adef.name}; use one of: {', '.join(allowed)}")
        return value
    if t == "int":
        try:
            return int(text.strip()) if text.strip() else ""
        except ValueError:
            raise ValueError(f"'{text}' is not a whole number") from None
    if t == "date":
        value = text.strip()
        if value:
            try:
                if not _DATE.match(value):
                    raise ValueError
                datetime.strptime(value, "%Y-%m-%d")
            except ValueError:
                raise ValueError(f"'{text}' is not a date written YYYY-MM-DD") from None
        return value
    if t == "uid-list":
        ids = _split(text)
        bad = [i for i in ids if not _UID.match(i)]
        if bad:
            raise ValueError(f"'{bad[0]}' is not an item ID such as SYS-0001")
        return ids
    if t == "string-list":
        return _split_escaped(text)
    if t == "ref-list":
        out = []
        for part in [p.strip() for p in text.split(";") if p.strip()]:
            if ":" not in part:
                raise ValueError(f"'{part}' must be written document-number:revision")
            docno, revision = part.rsplit(":", 1)
            out.append({"docno": docno.strip(), "revision": revision.strip()})
        return out
    raise ValueError(f"unsupported attribute type {t}")


def _same(new: Any, old: Any, adef: AttributeDef) -> bool:
    if adef.type in ("uid-list", "string-list"):
        return list(new) == [str(v) for v in (old or [])]
    if adef.type == "ref-list":
        return list(new) == [dict(v) for v in (old or [])]
    if adef.type == "int":
        return new == old or (new == "" and old in (None, ""))
    return str(new) == ("" if old is None else str(old))


def plan_import(
    cfg: ProjectConfig,
    items: Sequence[ItemData],
    rows: Sequence[Mapping[str, str]],
    *,
    why: str = "",
    root: Path | None = None,
) -> ImportPlan:
    """Plan (dry run) an import. With ``root`` the plan also knows which items are in a baseline, so a change that needs
    a reason is reported here and not half way through applying."""
    plan = ImportPlan()
    baselined = baselined_uids(root) if root is not None else frozenset()
    numbers: dict[str, int] = {}
    for existing_item in items:
        m = _UID.match(existing_item.uid)
        if m:
            numbers[existing_item.document] = max(
                numbers.get(existing_item.document, 0), int(existing_item.uid.rsplit("-", 1)[1])
            )
    known_columns = set(item_columns(cfg)) | HIDDEN_ATTRIBUTES
    existing = {i.uid: i for i in items}
    docs = {d.prefix: d for d in cfg.project.documents}
    defs = _attr_defs(cfg)
    reason_statuses = set(cfg.rules.get("change_control", {}).get("reason_required_statuses", []))

    all_columns = list(dict.fromkeys(c for r in rows for c in r if not c.startswith("_")))
    unknown = [c for c in all_columns if c not in known_columns]
    if unknown:
        plan.results.append(
            RowResult(
                1,
                "",
                "error",
                message=f"Unknown column(s): {', '.join(unknown)}. Use the column names of an exported file.",
            )
        )
        return plan

    seen: set[str] = set()
    pending_parent_checks: list[tuple[int, str, str, list[str]]] = []  # result index, doc, uid, parents

    def error(row_no: int, uid: str, message: str, doc: str = "") -> None:
        plan.results.append(RowResult(row_no, uid, "error", message=message, document=doc))

    for raw in rows:
        row_no = int(raw.get("_row", "0") or 0)
        uid = raw.get("id", "").strip()
        doc = raw.get("document", "").strip()
        item = existing.get(uid) if uid else None
        columns = [c for c in raw if not c.startswith("_")]  # each row (each sheet) is read with its own columns
        if raw.get("_error"):
            error(row_no, uid, f"{raw['_error']}.")
            continue
        if uid:
            if uid in seen:
                error(row_no, uid, f"{uid} appears twice in the file; keep one row per item.")
                continue
            seen.add(uid)
            prefix = uid.rsplit("-", 1)[0]
            if not _UID.match(uid):
                error(row_no, uid, f"'{uid}' is not an item ID such as SYS-0001.")
                continue
            if doc and doc != prefix:
                error(row_no, uid, f"{uid} belongs to document {prefix}, but the row says {doc}.")
                continue
            doc = prefix
            canonical = f"{prefix}{cfg.numbering.sep_for(prefix)}{int(uid.rsplit('-', 1)[1]):0{cfg.numbering.digits_for(prefix)}d}"
            if uid != canonical:
                error(
                    row_no,
                    uid,
                    f"Write the ID as {canonical} (with its leading zeros), or leave it empty for a new item.",
                )
                continue
            if item is None and int(uid.rsplit("-", 1)[1]) <= numbers.get(prefix, 0):
                error(
                    row_no,
                    uid,
                    f"{uid} is a new item but its number is not above the highest existing one in {prefix} "
                    f"({numbers[prefix]}). Use a higher number, or leave the ID empty.",
                )
                continue
        if not doc:
            error(row_no, uid, "The row needs a document (or an ID that names it).")
            continue
        decl = docs.get(doc)
        if decl is None:
            error(row_no, uid, f"Document {doc} does not exist in this project.", doc)
            continue
        kind_defs = defs[decl.kind]

        core: dict[str, Any] = {}
        attrs: dict[str, Any] = {}
        parents: list[str] | None = None
        problem = ""
        try:
            for col in columns:
                text = raw.get(col, "")
                if col in ("id", "document"):
                    continue
                if col == "level":
                    if text.strip() and not _LEVEL.match(text.strip()):
                        raise ValueError(f"'{text}' is not a level such as 1.2")
                    if text.strip():
                        core["level"] = text.strip()
                elif col in ("normative", "derived", "active"):
                    if text.strip():
                        core[col] = _bool(text)
                elif col in ("header", "ref"):
                    core[col] = text
                elif col == "text":
                    core["text"] = text.replace("\r\n", "\n").strip()
                elif col == "parents":
                    parents = [p for p in _split(text)]
                    bad = [p for p in parents if not _UID.match(p)]
                    if bad:
                        raise ValueError(f"'{bad[0]}' is not an item ID such as SYS-0001")
                elif col in HIDDEN_ATTRIBUTES:
                    continue
                else:
                    adef = kind_defs.get(col)
                    if adef is None:
                        if text.strip():
                            raise ValueError(f"'{col}' is not defined for {decl.kind} documents ({doc})")
                        continue
                    attrs[col] = _decode_attr(adef, text, cfg)
            for field_name, value in {**core, **attrs}.items():
                textcheck.check(field_name, value)
        except ValueError as exc:
            problem = str(exc)
        if problem:
            error(row_no, uid, f"{uid or 'New row'}: {problem}.", doc)
            continue

        if item is None:
            number = None
            if uid:
                number = int(uid.rsplit("-", 1)[1])
            for k in [k for k, v in attrs.items() if v in ("", [])]:
                del attrs[k]  # blank cells on new items keep the template defaults
            idx = len(plan.results)
            plan.results.append(
                RowResult(
                    row_no,
                    uid,
                    "create",
                    tuple(sorted([*core, *attrs, *(["parents"] if parents else [])])),
                    document=doc,
                )
            )
            plan.ops.append(_Op(idx, "create", doc, uid, number, core, attrs, parents, []))
            if parents:
                pending_parent_checks.append((idx, doc, uid, parents))
            continue

        changes: list[str] = []
        for key, new in core.items():
            old = {"text": item.text.strip(), "level": item.level, "normative": item.normative, "derived": item.derived,
                   "active": item.active, "header": item.header, "ref": item.ref}[key]  # fmt: skip
            if (str(new) if not isinstance(new, bool) else new) != (old if isinstance(old, bool) else str(old)):
                changes.append(key)
        changed_attrs = {}
        for key, new in attrs.items():
            if not _same(new, item.attrs.get(key), kind_defs[key]):
                changed_attrs[key] = new
                changes.append(key)
        if parents is not None and sorted(set(parents)) != sorted(item.links):
            changes.append("parents")
        if not changes:
            plan.results.append(RowResult(row_no, uid, "unchanged", document=doc))
            continue
        if (item.attrs.get("status") in reason_statuses or uid in baselined) and not why.strip():
            error(
                row_no,
                uid,
                f"{uid} is {item.attrs.get('status') if item.attrs.get('status') in reason_statuses else 'part of a baseline'}; "
                "changing it needs a reason. Enter why you are importing these changes.",
                doc,
            )
            continue
        idx = len(plan.results)
        plan.results.append(RowResult(row_no, uid, "update", tuple(sorted(changes)), document=doc))
        plan.ops.append(_Op(idx, "update", doc, uid, None, {k: v for k, v in core.items() if k in changes}, changed_attrs,
                            sorted(set(parents)) if "parents" in changes and parents is not None else None, sorted(changes)))  # fmt: skip
        if "parents" in changes and parents is not None:
            pending_parent_checks.append((idx, doc, uid, parents))

    # Parents must exist; a parent that is itself rejected does not count (--skip-errors must not create orphans).
    checked: set[int] = set()
    while True:
        known_uids = set(existing) | {
            op.uid
            for op in plan.ops
            if op.kind == "create" and op.uid and plan.results[op.result_index].action != "error"
        }
        failed = False
        for idx, doc, _uid, parents in pending_parent_checks:
            if idx in checked or plan.results[idx].action == "error":
                continue
            decl = docs[doc]
            problem = ""
            if parents and decl.parent is None:
                problem = f"{doc} is the root document and cannot have parents"
            else:
                for p in parents:
                    if p not in known_uids:
                        problem = f"parent {p} does not exist"
                    elif p.rsplit("-", 1)[0] != decl.parent:
                        problem = f"parent {p} is not in the parent document of {doc} ({decl.parent})"
                    if problem:
                        break
            if problem:
                old = plan.results[idx]
                plan.results[idx] = RowResult(
                    old.row, old.uid, "error", message=f"{old.uid or 'New row'}: {problem}.", document=doc
                )
                failed = True
            else:
                checked.add(idx)
        if not failed:
            break
    plan.ops = [op for op in plan.ops if plan.results[op.result_index].action != "error"]
    return plan


# applying ##############################################################################################
def apply_import(
    root: Path,
    plan: ImportPlan,
    *,
    user: str | None = None,
    why: str = "",
    skip_errors: bool = False,
    change_request: str | None = None,
) -> ImportReport:
    unchanged = plan.count("unchanged")
    if plan.errors and not skip_errors:
        return ImportReport(False, 0, 0, unchanged, len(plan.errors))
    svc = EditService(root, user=user, change_request=change_request)
    cfg = svc._cfg
    proj = DoorstopProject.open(root)
    order = {d.prefix: n for n, d in enumerate(cfg.project.documents)}
    created: list[tuple[_Op, str]] = []
    for op in (o for o in plan.ops if o.kind == "update"):  # every reason is checked before the first write
        svc.require_reason(proj.get_item(op.uid), why)
    # Creations first, ascending by explicit number (Doorstop cannot create a number below the next free one).
    creates = sorted(
        (o for o in plan.ops if o.kind == "create"),
        key=lambda o: (order[o.document], o.number is None, o.number or 0, o.result_index),
    )
    for op in creates:
        decl = cfg.project.document(op.document)
        assert decl is not None
        defaults = dict(cfg.templates.kinds[decl.kind].defaults)
        item = proj.add_item(
            op.document,
            op.core.get("text", ""),
            attrs={**defaults, **op.attrs},
            level=op.core.get("level"),
            normative=op.core.get("normative", True),
            derived=op.core.get("derived", False),
            header=op.core.get("header", ""),
            number=op.number,
            active=op.core.get("active", True),
            ref=op.core.get("ref", ""),
        )
        if op.uid and item.uid != op.uid:
            raise ProjectError(
                f"{op.uid} could not be created with that number (it became {item.uid}); the import stopped."
            )
        created.append((op, item.uid))
    updated = 0
    for op in (o for o in plan.ops if o.kind == "update"):
        svc.require_reason(proj.get_item(op.uid), why)
        proj.update_item(
            op.uid,
            text=op.core.get("text"),
            attrs=op.attrs or None,
            normative=op.core.get("normative"),
            derived=op.core.get("derived"),
            active=op.core.get("active"),
            header=op.core.get("header"),
            level=op.core.get("level"),
            ref=op.core.get("ref"),
        )
        if op.parents is not None:
            proj.set_links(op.uid, op.parents)
        svc.record(op.uid, "import", op.changes, why)
        updated += 1
    for op, uid in created:  # parents last: they may point at items created above
        if op.parents:
            proj.set_links(uid, op.parents)
        svc.record(uid, "import", ["text", *op.attrs, *(["parents"] if op.parents else [])], why)
    return ImportReport(True, len(created), updated, unchanged, len(plan.errors), tuple(u for _, u in created))
