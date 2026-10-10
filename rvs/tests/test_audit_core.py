"""Regressions for defects found by the full audit (core, CLI and exporters). One test per confirmed finding."""

import io
import json
import os
import stat
import sys
import time
from pathlib import Path

import pytest
import yaml
from dulwich.objects import Blob, Commit, Tree
from dulwich.repo import Repo

from rvs_core import atomicio
from rvs_core.adapter import DoorstopProject, ProjectError, UnreadableItemError
from rvs_core.authoring import EditService, history_path
from rvs_core.changecontrol import baselines as bl
from rvs_core.changecontrol.baselines import BaselineError, create_baseline, list_baselines, verify_baseline
from rvs_core.changecontrol.changes import ChangeRequestError, ChangeRequestStore
from rvs_core.config import load_project_config
from rvs_core.validate import validate_project
from rvs_core.vcs.git import GitRepo


def _codes(root: Path) -> set[str]:
    return {f.code for f in validate_project(root).findings}


# git: ignore rules, gitlinks, damaged metadata -------------------------------------------------------------------------------------
def test_gitignore_cannot_drop_project_data_from_a_baseline(git_project: Path):
    (git_project / ".gitignore").write_text("VER/\nbaselines/\nhistory/\n*.bin\n")
    (git_project / "attachment.bin").write_bytes(b"\0secret")
    baseline = create_baseline(git_project, "G1", "x", user="a")
    assert verify_baseline(git_project, "G1") == []
    repo = GitRepo.discover(git_project)
    paths = repo.tree_paths(baseline.commit, git_project)
    assert "baselines/G1.yaml" in paths and "VER/VER-0001.yml" in paths
    assert "attachment.bin" not in paths  # ignored attachments still follow the ignore rules


def test_a_commit_that_leaves_out_project_files_is_refused_and_taken_back(git_project: Path, monkeypatch):  # type: ignore[no-untyped-def]
    from rvs_core.vcs import git as gitmod

    monkeypatch.setattr(gitmod, "PROJECT_DATA_SUFFIXES", ())  # as if the data files followed the ignore rules
    (git_project / ".gitignore").write_text("VER/\n")
    EditService(git_project).update_item("SYS-0001", attrs={"status": "approved"})
    repo = GitRepo.discover(git_project)
    head_before = repo.head()
    item_before = (git_project / "SYS" / "SYS-0001.yml").read_bytes()
    history_before = history_path(git_project, "SYS-0001").read_bytes()
    with pytest.raises(BaselineError, match="leave out project files"):
        create_baseline(git_project, "G2", "x", user="a")
    assert repo.head() == head_before and repo.tag("rvs/baseline/G2") is None
    assert (git_project / "SYS" / "SYS-0001.yml").read_bytes() == item_before
    assert not (git_project / "baselines" / "G2.yaml").exists()
    assert history_path(git_project, "SYS-0001").read_bytes() == history_before  # no "baseline" entry left behind
    assert not (git_project / "baselines").exists()  # created by the attempt, removed with it
    monkeypatch.undo()
    (git_project / ".gitignore").unlink()
    create_baseline_ok = create_baseline(git_project, "G2", "x", user="a")  # the name was not burned
    assert create_baseline_ok.commit


def test_a_failed_tag_takes_the_commit_back(git_project: Path, monkeypatch):  # type: ignore[no-untyped-def]
    repo = GitRepo.discover(git_project)
    create_baseline(git_project, "T0", "x", user="a")
    head = repo.head()

    def boom(*_a: object, **_k: object) -> None:
        raise bl.GitError("simulated tag failure")

    monkeypatch.setattr(GitRepo, "create_tag", boom)
    with pytest.raises(bl.GitError):
        create_baseline(git_project, "T1", "x", user="a")
    assert GitRepo.discover(git_project).head() == head  # no orphan "Baseline T1" commit
    assert not (git_project / "baselines" / "T1.yaml").exists()


def test_history_of_the_promotion_is_part_of_the_baseline_commit(git_project: Path):
    EditService(git_project).update_item("SYS-0001", attrs={"status": "approved"})
    baseline = create_baseline(git_project, "H1", "x", user="a")
    assert "history/SYS/SYS-0001.jsonl" in GitRepo.discover(git_project).tree_paths(baseline.commit, git_project)


def test_one_unrestorable_file_does_not_stop_the_other_restores(git_project: Path, monkeypatch):  # type: ignore[no-untyped-def]
    a, b = git_project / "a.txt", git_project / "b.txt"
    a.write_text("changed")
    b.write_text("changed")
    real = Path.write_bytes

    def flaky(self: Path, data: bytes) -> int:
        if self == a:
            raise PermissionError("locked")
        return real(self, data)

    monkeypatch.setattr(Path, "write_bytes", flaky)
    repo = GitRepo.discover(git_project)
    bl._roll_back({a: b"orig", b: b"orig"}, git_project / "baselines" / "x.yaml", repo, git_project, "", None)
    assert b.read_text() == "orig"


def test_creating_the_same_baseline_twice_keeps_the_first_manifest(git_project: Path):
    create_baseline(git_project, "D1", "first", user="a")
    manifest = (git_project / "baselines" / "D1.yaml").read_bytes()
    with pytest.raises(BaselineError, match="already exists"):
        create_baseline(git_project, "D1", "second", user="b")
    assert (git_project / "baselines" / "D1.yaml").read_bytes() == manifest
    assert verify_baseline(git_project, "D1") == []


def test_a_reserved_name_is_not_overwritten_by_a_loser(git_project: Path):
    """Two creators racing for one name: the one that lost must not delete the winner's manifest on its way out."""
    (git_project / "baselines").mkdir()
    (git_project / "baselines" / "R1.yaml").write_text("held by someone else\n")
    with pytest.raises(BaselineError):
        create_baseline(git_project, "R1", "x", user="a")
    assert (git_project / "baselines" / "R1.yaml").read_text() == "held by someone else\n"


@pytest.mark.parametrize("name", ["CON", "nul", "com1", "LPT9.txt", "aux.x"])
def test_windows_device_names_are_refused(git_project: Path, name: str):
    with pytest.raises(BaselineError, match="Windows"):
        create_baseline(git_project, name, "x", user="a")


def test_names_differing_only_in_case_are_refused(git_project: Path):
    create_baseline(git_project, "Case1", "x", user="a")
    with pytest.raises(BaselineError, match="upper/lower case|already exists"):
        create_baseline(git_project, "case1", "x", user="a")


def test_a_folder_with_dots_in_its_name_can_be_baselined(tmp_path: Path):
    from rvs_core.examples.minimal import build_minimal_project

    repo_root = tmp_path / "repo"
    root = repo_root / "v1..2"
    build_minimal_project(root)
    GitRepo.init(repo_root)
    assert create_baseline(root, "B1", "x", user="a").tag.startswith("rvs/baseline/v1_")
    assert verify_baseline(root, "B1") == []


def test_list_baselines_follows_a_renamed_project_folder(tmp_path: Path):
    from rvs_core.examples.minimal import build_minimal_project

    repo_root = tmp_path / "repo"
    old = repo_root / "proj"
    build_minimal_project(old)
    GitRepo.init(repo_root)
    create_baseline(old, "B1", "x", user="a")
    new = repo_root / "renamed"
    old.rename(new)
    (found,) = list_baselines(new)
    assert found.tag and found.commit and found.manifest_sha256


def test_damaged_tag_objects_do_not_make_the_project_unopenable(git_project: Path):
    create_baseline(git_project, "O1", "x", user="a")
    repo = Repo(str(git_project))
    repo.refs[b"refs/tags/rvs/baseline/Ghost"] = b"1" * 40  # points at an object this clone does not have
    repo.close()
    report = validate_project(git_project)
    assert report.exit_code != 3 and "RVS-VALIDATION-FAILED" not in {f.code for f in report.findings}


def test_a_gitlink_in_a_tagged_tree_does_not_break_reading_it(git_project: Path):
    baseline = create_baseline(git_project, "L1", "x", user="a")
    repo = Repo(str(git_project))
    commit = repo[baseline.commit.encode()]
    assert isinstance(commit, Commit)
    root_tree = repo[commit.tree]
    assert isinstance(root_tree, Tree)
    blob = Blob.from_string(b"x")
    repo.object_store.add_object(blob)
    root_tree.add(b"sub", 0o160000, b"2" * 40)  # a submodule entry
    repo.object_store.add_object(root_tree)
    new = Commit()
    new.tree, new.parents = root_tree.id, [commit.id]
    new.author = new.committer = b"x <x@x>"
    new.author_time = new.commit_time = 1
    new.author_timezone = new.commit_timezone = 0
    new.message = b"with gitlink"
    repo.object_store.add_object(new)
    repo.close()
    files = GitRepo.discover(git_project).read_tree(new.id.decode(), git_project)
    assert "SYS/SYS-0001.yml" in files and not any(p.startswith("sub") for p in files)


def test_commit_directory_stays_fast_with_many_files(git_project: Path):
    for n in range(3000):
        (git_project / "VER" / f"x{n}.txt").write_text("x")
    started = time.monotonic()
    GitRepo.discover(git_project).commit_directory(git_project, "many", "a")
    assert (
        time.monotonic() - started < 20
    )  # the stale-file scan used to be quadratic (8,000 files: 3 s of pure set building)


# change requests ---------------------------------------------------------------------------------------------------------------------
def test_a_change_request_with_a_foreign_id_does_not_hang_creation(git_project: Path):
    cfg, _ = load_project_config(git_project)
    store = ChangeRequestStore(git_project, cfg)
    store.create("first", "", "a")
    (git_project / "changes" / "CR-0002.yaml").write_text(
        yaml.safe_dump({"rvs_schema_version": 1, "id": "CR-0001", "title": "odd", "status": "open"})
    )
    cr = store.create("third", "", "a")  # used to spin forever on the file-name collision
    assert cr.id == "CR-0003"


def test_change_requests_work_without_hard_links(git_project: Path, monkeypatch):  # type: ignore[no-untyped-def]
    cfg, _ = load_project_config(git_project)

    def no_links(*_a: object, **_k: object) -> None:
        raise PermissionError(1, "Operation not permitted")

    monkeypatch.setattr(os, "link", no_links)
    store = ChangeRequestStore(git_project, cfg)
    assert store.create("a", "", "x").id == "CR-0001"
    assert store.create("b", "", "x").id == "CR-0002"


def test_a_corrupt_change_request_is_a_plain_error_and_a_finding(git_project: Path):
    cfg, _ = load_project_config(git_project)
    store = ChangeRequestStore(git_project, cfg)
    store.create("a", "", "x")
    (git_project / "changes" / "CR-0001.yaml").write_text("title: [unclosed\n")
    with pytest.raises(ChangeRequestError, match="not valid YAML"):
        store.list()
    report = validate_project(git_project)
    assert "RVS-CR-INVALID" in {f.code for f in report.findings} and "RVS-VALIDATION-FAILED" not in {
        f.code for f in report.findings
    }


# manifests ---------------------------------------------------------------------------------------------------------------------------
def test_a_corrupt_manifest_is_reported_not_raised(git_project: Path):
    create_baseline(git_project, "M1", "x", user="a")
    (git_project / "baselines" / "M1.yaml").write_text("items: [unclosed\n")
    EditService(git_project).update_item(
        "VER-0001", attrs={"title": "edit"}, why="r"
    )  # must not raise yaml.ParserError
    report = validate_project(git_project)
    assert "RVS-BASELINE-MODIFIED" in {f.code for f in report.findings}


# adapter ---------------------------------------------------------------------------------------------------------------------------
def test_include_cannot_be_smuggled_in_through_a_tag_handle(minimal_project: Path):
    secret = minimal_project.parent / "secret.txt"
    secret.write_text("TOPSECRET123")
    config = minimal_project / "SYS" / ".doorstop.yml"
    original = config.read_text()
    config.write_text("%TAG !x! !inc\n---\n" + original.replace("sep: '-'", "sep: !x!lude ../../secret.txt", 1))
    with pytest.raises(ProjectError, match="custom YAML tag"):
        DoorstopProject.open(minimal_project)
    report = validate_project(minimal_project)
    assert "TOPSECRET123" not in "".join(f.message for f in report.findings)


def test_an_alias_bomb_in_an_item_is_refused_quickly(minimal_project: Path):
    path = minimal_project / "SYS" / "SYS-0001.yml"
    lines = ["bomb0: &a0 [x, x, x, x, x, x, x, x]"] + [
        f"bomb{i}: &a{i} [{', '.join([f'*a{i - 1}'] * 8)}]" for i in range(1, 10)
    ]
    path.write_text(path.read_text() + "\n".join(lines) + "\n")
    time.sleep(0.01)
    started = time.monotonic()
    with pytest.raises(UnreadableItemError, match="aliases"):
        DoorstopProject.open(minimal_project).items()
    assert time.monotonic() - started < 10


@pytest.mark.skipif(sys.platform == "win32", reason="POSIX permissions and symbolic links")
def test_saving_an_item_keeps_permissions_and_symlinks(minimal_project: Path):
    path = minimal_project / "SYS" / "SYS-0001.yml"
    path.chmod(0o660)
    EditService(minimal_project).update_item("SYS-0001", attrs={"owner": "someone"})
    assert stat.S_IMODE(path.stat().st_mode) == 0o660
    real = minimal_project / "elsewhere.yml"
    real.write_bytes(path.read_bytes())
    path.unlink()
    path.symlink_to(real)
    EditService(minimal_project).update_item("SYS-0001", attrs={"owner": "another"})
    assert path.is_symlink() and "another" in real.read_text()


def test_all_uids_include_inactive_items(minimal_project: Path):
    proj = DoorstopProject.open(minimal_project)
    proj.update_item("VER-0001", active=False)
    proj = DoorstopProject.open(minimal_project)
    assert "VER-0001" in proj.all_uids() and "VER-0001" not in {i.uid for i in proj.items()}


# import ----------------------------------------------------------------------------------------------------------------------------
def test_an_inactive_item_blocks_a_create_row_with_its_id(minimal_project: Path):
    from rvs_core.exporters.itemsio import plan_import

    proj = DoorstopProject.open(minimal_project)
    last = max(i.uid for i in proj.items() if i.document == "VER")
    proj.update_item(last, active=False)
    proj = DoorstopProject.open(minimal_project)
    cfg, _ = load_project_config(minimal_project)
    rows = [{"id": last, "document": "VER", "text": "back again", "_row": "2"}]
    plan = plan_import(cfg, proj.items(), rows, root=minimal_project)
    assert plan.errors and "inactive" in plan.errors[0].message
    new_number = int(last.rsplit("-", 1)[1]) + 1
    assert not plan_import(
        cfg,
        proj.items(),
        [{"id": f"VER-{new_number:04d}", "document": "VER", "text": "t", "_row": "2"}],
        root=minimal_project,
    ).errors


def test_discard_item_removes_the_file(minimal_project: Path):
    proj = DoorstopProject.open(minimal_project)
    item = proj.add_item("VER", "stray")
    path = minimal_project / item.path
    assert path.exists()
    proj.discard_item(item.uid)
    assert not path.exists()


def test_xlsx_rows_after_a_long_gap_are_still_read():
    from openpyxl import Workbook

    from rvs_core.exporters.itemsio import read_xlsx

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["id", "text"])
    ws.append(["SYS-0001", "first"])
    ws["A3000"], ws["B3000"] = "SYS-0002", "after a gap of 2,997 empty rows"
    buf = io.BytesIO()
    wb.save(buf)
    assert [r["id"] for r in read_xlsx(buf.getvalue())] == ["SYS-0001", "SYS-0002"]


@pytest.mark.parametrize("encoding", ["utf-16", "utf-16-le", "utf-16-be", "utf-32-be", "utf-8"])
def test_reqif_doctype_is_refused_in_every_encoding(encoding: str):
    from rvs_core.exporters import reqif

    text = '<?xml version="1.0"?><!DOCTYPE REQ-IF [<!ENTITY a "aaaa">]><REQ-IF xmlns="http://www.omg.org/spec/ReqIF/20110401/reqif.xsd"/>'
    data = text.encode(encoding)
    if encoding in ("utf-16-le", "utf-16-be", "utf-32-be"):
        data = text.replace('<?xml version="1.0"?>', "").encode(encoding)  # no declaration, no byte-order mark
    with pytest.raises(ValueError, match="DOCTYPE|not a readable"):
        reqif._parse(data)


def test_type_checks_apply_to_edits(minimal_project: Path):
    svc = EditService(minimal_project)
    with pytest.raises(ValueError, match="executed_on"):
        svc.update_item("VER-0001", attrs={"executed_on": "next tuesday"})
    with pytest.raises(ValueError, match="link_verifies"):
        svc.update_item("VER-0001", attrs={"link_verifies": ["NOT-AN-ID"]})


def test_item_ids_with_a_trailing_newline_or_foreign_digits_are_not_ids():
    from rvs_core.attrtypes import UID

    assert UID.match("SYS-0001") and not UID.match("SYS-0001\n") and not UID.match("SYS-١٢")


# traceability ----------------------------------------------------------------------------------------------------------------------
def test_impact_of_a_very_deep_chain_does_not_overflow():
    from rvs_core.trace.impact import ImpactNode, ImpactResult

    root = node = ImpactNode("A-0")
    for n in range(1, 3000):
        child = ImpactNode(f"A-{n}", n, "parent")
        node.children.append(child)
        node = child
    assert ImpactResult(root).count == 2999


# outputs ---------------------------------------------------------------------------------------------------------------------------
def test_source_date_epoch_pins_the_generation_time(minimal_project: Path, monkeypatch):  # type: ignore[no-untyped-def]
    from rvs_core.matrices.provenance import Provenance

    cfg, _ = load_project_config(minimal_project)
    monkeypatch.setenv("SOURCE_DATE_EPOCH", "1700000000")
    a = Provenance.now(cfg, user="u")
    time.sleep(1.1)
    assert a.generated == Provenance.now(cfg, user="u").generated
    assert a.generated.year == 2023


def test_xlsx_export_writes_no_temporary_files(minimal_project: Path, tmp_path: Path, monkeypatch):  # type: ignore[no-untyped-def]
    import tempfile

    from rvs_core.exporters.export_request import ExportRequest, build_output
    from rvs_core.matrices.provenance import Provenance

    scratch = tmp_path / "tmp"
    scratch.mkdir()
    monkeypatch.setattr(tempfile, "tempdir", str(scratch))
    report = validate_project(minimal_project, doorstop=False)
    assert report.config is not None and report.graph is not None
    prov = Provenance.now(report.config, user="u")
    for kind in ("items", "vcm"):
        data = build_output(ExportRequest(kind, "xlsx"), report.config, report.items, report.graph, prov)
        assert data[:2] == b"PK"
    assert list(scratch.iterdir()) == []


def test_matrix_headers_cannot_carry_formulas():
    from datetime import datetime

    from rvs_core.matrices.provenance import Provenance
    from rvs_core.matrices.render import to_csv
    from rvs_core.matrices.table import MatrixTable

    prov = Provenance("1", "3.2", "p", "working copy", datetime(2026, 1, 1), "u")
    table = MatrixTable("t", ["=HYPERLINK(1)", "ok"], [["a", "b"]], [None], prov)
    assert any(line.startswith("'=HYPERLINK") for line in to_csv(table).splitlines())


def test_a_long_statement_does_not_make_pdf_export_explode():
    from datetime import datetime

    from rvs_core.exporters.model import Doc, Entry
    from rvs_core.exporters.pdf_out import render_pdf
    from rvs_core.matrices.provenance import Provenance

    prov = Provenance("1", "3.2", "p", "working copy", datetime(2026, 1, 1), "u")
    doc = Doc("t", prov, [Entry("SYS-0001", "title", "word " * 40_000, (), "")])
    started = time.monotonic()
    assert render_pdf(doc)[:4] == b"%PDF"
    assert time.monotonic() - started < 30  # 200 KB took two minutes; 5 MB never finished


def test_docx_body_keeps_section_properties_last():
    from docx import Document

    from rvs_core.exporters import docx_out  # noqa: F401 - applies the patch

    d = Document()
    d.add_paragraph("one")
    d.add_table(rows=1, cols=1)
    d.add_paragraph("two")
    assert [c.tag.split("}")[1] for c in d.element.body] == ["p", "tbl", "p", "sectPr"]


def test_a_failed_atomic_write_keeps_the_old_file(tmp_path: Path, monkeypatch):  # type: ignore[no-untyped-def]
    target = tmp_path / "out.csv"
    target.write_bytes(b"old")
    monkeypatch.setattr(os, "replace", lambda *_a, **_k: (_ for _ in ()).throw(OSError(28, "No space left on device")))
    with pytest.raises(OSError):
        atomicio.write_bytes(target, b"new data")
    assert target.read_bytes() == b"old" and list(tmp_path.iterdir()) == [target]


# cache -----------------------------------------------------------------------------------------------------------------------------
def test_a_cache_that_this_computer_did_not_write_is_ignored(minimal_project: Path):
    DoorstopProject.open(minimal_project).items()
    time.sleep(2.2)  # old enough to be cached
    for p in (minimal_project / "SYS").glob("*.yml"):
        os.utime(p, (1, 1))
    proj = DoorstopProject.open(minimal_project)
    proj.items()
    cache_file = minimal_project / ".rvs-cache" / "SYS.json"
    assert cache_file.is_file()
    forged = json.loads(cache_file.read_text())
    first = next(iter(forged["items"].values()))
    first["d"]["text"] = "shall NOT provide power"
    cache_file.write_text(json.dumps(forged))  # the signature no longer matches
    assert "shall NOT provide power" not in {i.text for i in DoorstopProject.open(minimal_project).items()}
    forged["sig"] = "0" * 64
    cache_file.write_text(json.dumps(forged))
    assert "shall NOT provide power" not in {i.text for i in DoorstopProject.open(minimal_project).items()}


def test_an_unsigned_snapshot_folder_is_extracted_again(git_project: Path):
    create_baseline(git_project, "S1", "x", user="a")
    snap = bl.snapshot_dir(git_project, "S1")
    (snap / "SYS" / "SYS-0001.yml").write_text("forged: true\n")
    (snap / ".complete").write_text("")
    again = bl.snapshot_dir(git_project, "S1")
    assert "forged" not in (again / "SYS" / "SYS-0001.yml").read_text()


# CLI -------------------------------------------------------------------------------------------------------------------------------
def test_the_command_line_never_prints_a_traceback(minimal_project: Path, monkeypatch, capsys):  # type: ignore[no-untyped-def]
    import rvs_cli.main as cli

    def boom(*_a: object, **_k: object) -> int:
        raise RuntimeError("secret statement text")

    monkeypatch.setattr(cli, "_validate", boom)
    monkeypatch.delenv("RVS_DEBUG", raising=False)
    assert cli.main(["validate", str(minimal_project)]) == 1
    err = capsys.readouterr().err
    assert "unexpected error (RuntimeError)" in err and "Traceback" not in err and "secret statement text" not in err
    monkeypatch.setattr(cli, "_validate", lambda *_a, **_k: (_ for _ in ()).throw(KeyboardInterrupt()))
    assert cli.main(["validate", str(minimal_project)]) == 130


def test_diff_on_an_unreadable_item_is_a_message(git_project: Path, capsys):  # type: ignore[no-untyped-def]
    import rvs_cli.main as cli

    (git_project / "SYS" / "SYS-0001.yml").write_text("text: [unclosed\n")
    assert cli.main(["baseline", "create", str(git_project), "U1", "-m", "x"]) in (2, 3)
    err = capsys.readouterr().err
    assert "Traceback" not in err and err.strip()


def test_validation_safety_net_saves_a_crash_report(minimal_project: Path, monkeypatch):  # type: ignore[no-untyped-def]
    import rvs_core.validate as v

    monkeypatch.delenv("RVS_DEBUG", raising=False)
    monkeypatch.setattr(v, "_validate", lambda *_a, **_k: (_ for _ in ()).throw(AttributeError("boom")))
    (finding,) = v.validate_project(minimal_project).findings
    assert (
        finding.code == "RVS-VALIDATION-FAILED" and "crash" in finding.message.lower() and "saved to" in finding.message
    )
    monkeypatch.setenv("RVS_DEBUG", "1")
    with pytest.raises(AttributeError):
        v.validate_project(minimal_project)


def test_python_version_is_supported():
    assert sys.version_info >= (3, 11)


def test_export_format_follows_the_output_file_name(minimal_project: Path, tmp_path: Path):
    import rvs_cli.main as cli

    out = tmp_path / "vcm.xlsx"
    assert cli.main(["export", str(minimal_project), "--vcm", "-o", str(out)]) == 0
    assert out.read_bytes()[:2] == b"PK"  # a real workbook, not CSV text in a file called .xlsx
    spec = tmp_path / "spec.docx"
    assert cli.main(["export", str(minimal_project), "--spec", "-o", str(spec)]) == 0
    assert spec.read_bytes()[:2] == b"PK"
    assert cli.main(["export", str(minimal_project), "--spec", "-o", str(tmp_path / "s.html")]) == 0


def test_project_files_are_written_with_unix_line_endings_on_every_system():
    """Doorstop defaults to the operating system's line ending (CRLF on Windows), which made item files differ between
    systems and from the committed examples; the adapter pins it."""
    from doorstop import settings

    import rvs_core.adapter  # noqa: F401 - applies the settings

    assert settings.WRITE_LINESEPERATOR == "\n"


def test_paths_are_posix_and_zip_members_do_not_depend_on_the_platform(minimal_project: Path):
    """On Windows item paths came back with backslashes and ZIP members carried a different 'created on' byte, which broke
    baselines (tree paths use '/') and made DOCX/XLSX outputs differ between systems."""
    import zipfile

    from rvs_core.exporters.zipnorm import normalize_zip

    assert all("\\" not in i.path for i in DoorstopProject.open(minimal_project).items())
    raw = io.BytesIO()
    with zipfile.ZipFile(raw, "w") as z:
        z.writestr("a.txt", "x")
    with zipfile.ZipFile(io.BytesIO(normalize_zip(raw.getvalue()))) as z:
        assert {i.create_system for i in z.infolist()} == {3}
