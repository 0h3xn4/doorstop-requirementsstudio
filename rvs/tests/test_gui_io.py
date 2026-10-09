"""pytest-qt tests for M4: background export (responsive window), export and import dialogs."""

import shutil
import time
from pathlib import Path

import pytest
from PySide6.QtCore import QTimer

from conftest import EXAMPLES
from rvs_core.exporters.export_request import ExportRequest
from rvs_core.exporters.itemsio import read_csv
from rvs_gui.app import MainWindow, create_main_window


@pytest.fixture
def win(qtbot, minimal_project: Path) -> MainWindow:  # type: ignore[no-untyped-def]
    w = create_main_window()
    qtbot.addWidget(w)
    w.show()
    assert w.open_project(minimal_project)
    return w


def test_menu_has_import_and_export_and_nothing_online(win: MainWindow):
    from PySide6.QtGui import QAction

    texts = [a.text() for a in win.findChildren(QAction)]
    assert any("Export" in t for t in texts) and any("Import" in t for t in texts)
    assert not [t for t in texts if "online" in t.lower() or "update" in t.lower()]


# export ---------------------------------------------------------------------------------------------
def test_export_runs_in_background_and_reports_success(win: MainWindow, qtbot, tmp_path: Path):  # type: ignore[no-untyped-def]
    from pypdf import PdfReader

    target = tmp_path / "vcm.pdf"
    with qtbot.waitSignal(win.export_done, timeout=60000) as blocker:
        win.export_to(ExportRequest("vcm", "pdf"), target)
    assert blocker.args == [str(target), ""]
    assert PdfReader(str(target)).pages and win.notification.kind == "success"
    assert "vcm.pdf" in win.notification.text()


@pytest.mark.parametrize(("kind", "fmt"), [("items", "xlsx"), ("spec", "docx"), ("coverage", "html"), ("vcm", "csv")])
def test_export_kinds(win: MainWindow, qtbot, tmp_path: Path, kind: str, fmt: str):  # type: ignore[no-untyped-def]
    target = tmp_path / f"out.{fmt}"
    with qtbot.waitSignal(win.export_done, timeout=60000):
        win.export_to(ExportRequest(kind, fmt), target)
    assert target.stat().st_size > 100


def test_export_failure_is_a_friendly_message(win: MainWindow, qtbot, tmp_path: Path):  # type: ignore[no-untyped-def]
    target = tmp_path / "no-such-folder" / "x.pdf"
    with qtbot.waitSignal(win.export_done, timeout=60000) as blocker:
        win.export_to(ExportRequest("vcm", "pdf"), target)
    assert blocker.args[1] and win.notification.kind == "error"
    text = win.notification.text()
    assert "x.pdf" in text and "Traceback" not in text and "Choose another" in text


def test_invalid_request_is_reported(win: MainWindow, qtbot, tmp_path: Path):  # type: ignore[no-untyped-def]
    with qtbot.waitSignal(win.export_done, timeout=60000):
        win.export_to(ExportRequest("items", "pdf"), tmp_path / "x.pdf")
    assert win.notification.kind == "error" and "csv, xlsx" in win.notification.text()


def test_window_stays_responsive_while_exporting(qtbot, tmp_path: Path):  # type: ignore[no-untyped-def]
    root = tmp_path / "sat"
    shutil.copytree(EXAMPLES / "satellite300", root)
    w = create_main_window()
    qtbot.addWidget(w)
    w.show()
    assert w.open_project(root)
    ticks: list[float] = []
    timer = QTimer()
    timer.setInterval(20)
    timer.timeout.connect(lambda: ticks.append(time.perf_counter()))
    timer.start()
    started = time.perf_counter()
    with qtbot.waitSignal(w.export_done, timeout=120000):
        w.export_to(ExportRequest("spec", "pdf"), tmp_path / "spec.pdf")
    timer.stop()
    duration = time.perf_counter() - started
    assert (tmp_path / "spec.pdf").stat().st_size > 10_000
    # the event loop kept running during the export: timer ticks arrived (at least one per ~100 ms), and no gap was long
    assert len(ticks) >= max(3, int(duration * 10))
    assert max(b - a for a, b in zip(ticks, ticks[1:], strict=False)) < 0.5


def test_export_dialog_follows_the_chosen_content(win: MainWindow, qtbot, tmp_path: Path):  # type: ignore[no-untyped-def]
    from rvs_gui.export_dialog import ExportDialog

    dlg = ExportDialog(win.session, current_uid="SYS-0002")
    qtbot.addWidget(dlg)
    dlg.set_kind("items")
    assert [dlg.format.itemText(i) for i in range(dlg.format.count())] == ["csv", "xlsx"]
    dlg.set_kind("spec")
    assert [dlg.format.itemText(i) for i in range(dlg.format.count())] == ["html", "docx", "pdf"]
    dlg.set_kind("trace")
    dlg.trace_source.setCurrentText("SYS")
    dlg.trace_target.setCurrentText("EPS")
    dlg.trace_direction.setCurrentText("up")
    dlg.format.setCurrentText("xlsx")
    dlg.path.setText(str(tmp_path / "t"))
    req, path = dlg.request(), dlg.destination()
    assert req == ExportRequest("trace", "xlsx", trace=("SYS", "EPS", "up"))
    assert path == tmp_path / "t.xlsx"  # the extension follows the format
    dlg.set_kind("impact")
    assert dlg.request().impact_uid == "SYS-0002"


def test_matrix_tab_exports_by_file_suffix(win: MainWindow, qtbot, tmp_path: Path):  # type: ignore[no-untyped-def]
    from openpyxl import load_workbook

    target = tmp_path / "vcm.xlsx"
    with qtbot.waitSignal(win.vcm_view.export_done, timeout=60000):
        win.vcm_view.export_file(target)
    assert load_workbook(target)["Matrix"].max_row == 7
    bad = tmp_path / "vcm.rtf"
    win.vcm_view.export_file(bad)
    assert win.notification.kind == "error" and "rtf" in win.notification.text() and not bad.exists()


# import ---------------------------------------------------------------------------------------------------
def _export_items(win: MainWindow, qtbot, path: Path) -> None:  # type: ignore[no-untyped-def]
    with qtbot.waitSignal(win.export_done, timeout=60000):
        win.export_to(ExportRequest("items", path.suffix.lstrip(".")), path)


def test_import_preview_and_apply(win: MainWindow, qtbot, tmp_path: Path):  # type: ignore[no-untyped-def]
    from rvs_gui.import_dialog import ImportDialog

    csv = tmp_path / "items.csv"
    _export_items(win, qtbot, csv)
    csv.write_text(csv.read_text("utf-8-sig").replace("Payload power", "Payload power v2"), encoding="utf-8")
    rows = win.read_import_file(csv)
    assert rows is not None and len(rows) == 10
    dlg = ImportDialog(win.session, rows)
    qtbot.addWidget(dlg)
    assert "1 to update" in dlg.summary.text() and dlg.apply_button.isEnabled()
    assert [dlg.model.index(r, 1).data() for r in range(dlg.model.rowCount())] == ["update"]
    dlg.reason.setText("customer review")
    win.apply_import(dlg.plan, why=dlg.reason.text(), skip_errors=False)
    assert win.session.item("SYS-0001").attrs["title"] == "Payload power v2"  # type: ignore[union-attr]
    assert win.notification.kind == "success" and "1 updated" in win.notification.text()
    assert win.table_model.index(win.table_model.row_of("SYS-0001"), 2).data() == "Payload power v2"


def test_import_dialog_blocks_errors_unless_skipped(win: MainWindow, qtbot, tmp_path: Path):  # type: ignore[no-untyped-def]
    from rvs_gui.import_dialog import ImportDialog

    rows = read_csv(b"id,status,title\nSYS-0001,finished,x\nSYS-0002,,Fine title\n")
    dlg = ImportDialog(win.session, rows)
    qtbot.addWidget(dlg)
    assert "1 error" in dlg.summary.text() and not dlg.apply_button.isEnabled()
    dlg.skip_errors.setChecked(True)
    assert dlg.apply_button.isEnabled()


def test_import_dialog_replans_when_the_reason_is_entered(win: MainWindow, qtbot):  # type: ignore[no-untyped-def]
    from rvs_gui.import_dialog import ImportDialog

    win.session.update_item("SYS-0001", attrs={"status": "baselined"})
    dlg = ImportDialog(win.session, read_csv(b"id,title\nSYS-0001,Changed\n"))
    qtbot.addWidget(dlg)
    assert "reason" in dlg.model.index(0, 3).data().lower() and not dlg.apply_button.isEnabled()
    dlg.reason.setText("customer request")
    assert dlg.apply_button.isEnabled() and dlg.model.index(0, 1).data() == "update"


def test_import_xlsx_file_is_read(win: MainWindow, qtbot, tmp_path: Path):  # type: ignore[no-untyped-def]
    xlsx = tmp_path / "items.xlsx"
    _export_items(win, qtbot, xlsx)
    rows = win.read_import_file(xlsx)
    assert rows is not None and len(rows) == 10


def test_import_wrong_file_type_or_corrupt_file_is_a_message(win: MainWindow, tmp_path: Path):
    odd = tmp_path / "items.txt"
    odd.write_text("x")
    assert (
        win.read_import_file(odd) is None
        and win.notification.kind == "error"
        and ".csv or .xlsx" in win.notification.text()
    )
    broken = tmp_path / "broken.xlsx"
    broken.write_bytes(b"not a zip")
    assert win.read_import_file(broken) is None and "broken.xlsx" in win.notification.text()
    assert "Traceback" not in win.notification.text()
