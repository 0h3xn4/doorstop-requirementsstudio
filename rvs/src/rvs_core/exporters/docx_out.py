"""DOCX via python-docx (bundled, offline). Fonts are referenced by name (IBM Plex Sans); Word falls back if absent."""

import io

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph as DocxParagraph

from rvs_core.exporters.mdlite import Bullets, Spans, parse
from rvs_core.exporters.model import Doc, Entry, Heading, Paragraph, TableBlock
from rvs_core.exporters.zipnorm import normalize_zip

SANS, MONO = "IBM Plex Sans", "IBM Plex Mono"
_GAP_FILL = {"unverified-approved": "FFD7D9", "orphan": "FFD7D9", "unverified": "FCF4D6", "childless": "FCF4D6"}


def _font(style_or_run, name: str, size: float | None = None) -> None:  # type: ignore[no-untyped-def]
    style_or_run.font.name = name
    r_pr = style_or_run.element.get_or_add_rPr() if hasattr(style_or_run.element, "get_or_add_rPr") else None
    if r_pr is not None:
        fonts = r_pr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            r_pr.append(fonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            fonts.set(qn(attr), name)
    if size:
        style_or_run.font.size = Pt(size)


def _shade(cell, fill: str) -> None:  # type: ignore[no-untyped-def]
    pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def _field(paragraph: DocxParagraph, instr: str) -> None:
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, instr), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def _add_spans(paragraph: DocxParagraph, spans: Spans, size: float | None = None) -> None:
    for text, style in spans:
        run = paragraph.add_run(text)
        run.bold = style == "b" or None
        run.italic = style == "i" or None
        if style == "code":
            _font(run, MONO, size)
        elif size:
            run.font.size = Pt(size)


def _markdown(document, text: str) -> None:  # type: ignore[no-untyped-def]
    for block in parse(text):
        if isinstance(block, Bullets):
            for item in block.items:
                _add_spans(document.add_paragraph(style="List Number" if block.ordered else "List Bullet"), item)
        else:
            _add_spans(document.add_paragraph(), block.spans)


def _table(document, t: TableBlock) -> None:  # type: ignore[no-untyped-def]
    table = document.add_table(rows=1, cols=len(t.columns))
    table.style = "Table Grid"
    for cell, name in zip(table.rows[0].cells, t.columns, strict=True):
        cell.text = ""
        run = cell.paragraphs[0].add_run(name)
        run.bold = True
        run.font.size = Pt(9)
        _shade(cell, "F4F4F4")
    header_pr = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_pr.append(repeat)
    for n, row in enumerate(t.rows):
        cells = table.add_row().cells
        flag = t.flags[n] if n < len(t.flags) else None
        for cell, value in zip(cells, row, strict=True):
            cell.text = ""
            cell.paragraphs[0].add_run(value).font.size = Pt(9)
            if flag:
                _shade(cell, _GAP_FILL.get(flag, "FCF4D6"))


def render_docx(doc: Doc) -> bytes:
    d = Document()
    _font(d.styles["Normal"], SANS, 10)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3", "List Bullet", "List Number"):
        _font(d.styles[name], SANS)
    section = d.sections[0]
    if doc.landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = section.right_margin = Cm(1.8)
    section.top_margin = section.bottom_margin = Cm(2.0)

    header = section.header.paragraphs[0]
    header.text = f"{doc.provenance.project} · {doc.provenance.baseline} · generated "
    header.add_run(f"{doc.provenance.generated.strftime('%Y-%m-%d %H:%M:%S')} by {doc.provenance.user}")
    for run in header.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x52, 0x52, 0x52)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Page ")
    _field(footer, "PAGE")
    footer.add_run(" of ")
    _field(footer, "NUMPAGES")

    d.add_paragraph(doc.title, style="Title")
    for line in doc.provenance.lines():
        p = d.add_paragraph()
        run = p.add_run(line)
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x52, 0x52, 0x52)
        p.paragraph_format.space_after = Pt(0)
    for note in doc.notes:
        p = d.add_paragraph()
        p.add_run(note).italic = True

    for block in doc.blocks:
        if isinstance(block, Heading):
            d.add_heading(block.text, level=min(block.level, 3))
        elif isinstance(block, Paragraph):
            _markdown(d, block.text)
        elif isinstance(block, TableBlock):
            _table(d, block)
        elif isinstance(block, Entry):
            d.add_heading(f"{block.uid} — {block.title}" if block.title else block.uid, level=3)
            _markdown(d, block.text)
            if block.fields:  # one paragraph per entry: far fewer document objects than one per field
                p = d.add_paragraph()
                p.paragraph_format.space_after = Pt(2)
                for n, (key, value) in enumerate(block.fields):
                    if n:
                        p.add_run("  ·  ").font.size = Pt(9)
                    k = p.add_run(f"{key}: ")
                    k.bold = True
                    k.font.size = Pt(9)
                    p.add_run(value).font.size = Pt(9)
            if block.rationale:
                p = d.add_paragraph()
                p.add_run("Rationale: ").bold = True
                p.add_run(block.rationale)

    props = d.core_properties
    props.title, props.author, props.last_modified_by = doc.title, doc.provenance.user, doc.provenance.user
    props.subject, props.comments = doc.provenance.project, f"Generated by rvs {doc.provenance.tool_version}"
    props.created = props.modified = doc.provenance.generated
    props.revision = 1
    buf = io.BytesIO()
    d.save(buf)
    return normalize_zip(buf.getvalue())
